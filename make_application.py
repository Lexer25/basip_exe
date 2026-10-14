from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

FONT = "Times New Roman"

doc = Document()

# --- page setup: A4 with standard document margins ---
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(3)
section.right_margin = Cm(1.5)

# --- base style ---
style = doc.styles["Normal"]
style.font.name = FONT
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
style.element.rPr.rFonts.set(qn("w:cs"), FONT)
pf = style.paragraph_format
pf.space_before = Pt(0)
pf.space_after = Pt(0)
pf.line_spacing_rule = WD_LINE_SPACING.SINGLE


def para(text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False,
         first_line=None, space_after=0, size=12):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if first_line is not None:
        p.paragraph_format.first_line_indent = Cm(first_line)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = FONT
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run._element.rPr.rFonts.set(qn("w:cs"), FONT)
    return p


# --- header block, right-aligned, 2 cm inset from the right margin ---
def header_line(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.right_indent = Cm(2)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(12)
    return p


header_line("Генеральному директору")
header_line("ООО «РиК»")
header_line("Смирнову А. В.")
header_line("от Иванова Петра Семеновича,")
header_line("проживающего по адресу:")
header_line("г. Москва, ул. Садовая, д. 10, кв. 25,")
header_line("тел. +7 (900) 123-45-67")
para(space_after=12)

# --- title ---
para("ЗАЯВЛЕНИЕ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_after=12)

# --- body ---
para("Прошу принять меня на работу в ООО «РиК» на должность менеджера "
     "по продажам с 01 сентября 2025 года.",
     first_line=1.25, space_after=12)

para("Квалификационные требованиям, предъявляемым к данной должности, "
     "соответствую. Трудовой договор прошу заключить на неопределенный срок "
     "с установлением должностного оклада согласно штатному расписанию.",
     first_line=1.25, space_after=24)

# --- signature block ---
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(24)
p.add_run("«01» августа 2025 г.")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p.paragraph_format.right_indent = Cm(1.5)
p.add_run("____________________ / П. С. Иванов")

doc.save("Заявление_о_приеме_на_работу_Иванов.docx")
print("saved")
